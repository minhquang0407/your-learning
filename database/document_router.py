"""Multimodal Document Router & Ingestion Engine.

Routes documents by format (.pdf, .pptx, .docx, .md, .txt) and document type
(PAPER, BOOK, SLIDE), utilizing Gemini Multimodal Vision, deterministic page bounds
validation, and fallback to local typographic processors.
"""

import io
import logging
import os
from typing import Any, Dict, List, Optional, Tuple, Union

from core.schemas import DocumentStructure, DocumentType, SubSectionNode
from database.document_processor import MathAwareDocumentProcessor, sub_chunk_section
from database.gemini_extractor import GeminiStructureExtractor
from database.page_validator import PageBoundsValidator

logger = logging.getLogger(__name__)


class DocumentRouter:
    """Central traffic controller for document processing and hierarchical chunking."""

    def __init__(
        self,
        gemini_extractor: Optional[GeminiStructureExtractor] = None,
        page_validator: Optional[PageBoundsValidator] = None,
        fallback_processor: Optional[MathAwareDocumentProcessor] = None,
    ):
        """Initialize the router with extractor, page validator, and fallback processor."""
        self.gemini_extractor = gemini_extractor or GeminiStructureExtractor()
        self.page_validator = page_validator or PageBoundsValidator()
        self.fallback_processor = fallback_processor or MathAwareDocumentProcessor()

    def process_document(
        self,
        content_or_bytes: Union[str, bytes, None] = None,
        file_path: Optional[str] = None,
        file_bytes: Optional[bytes] = None,
        filename: str = "",
        content: Optional[str] = None,
    ) -> Tuple[List[Dict[str, Any]], Dict[str, List[str]]]:
        """Route and process a document into pedagogical chunks and Table of Contents.

        Supports:
        - PPTX: python-pptx slide-by-slide + speaker notes.
        - DOCX: python-docx native headings hierarchy.
        - PDF: Gemini Multimodal Vision + PageBoundsValidator -> Type-specific slicing (SLIDE, PAPER, BOOK).
        - Markdown/Text/Fallbacks: Local typographic processor.

        Returns:
            Tuple containing:
            - final_documents: List of dicts with 'page_content' and 'metadata'.
            - toc_tree: Dict mapping parent headers to lists of sections.
        """
        resolved_filename = filename or ""
        if not resolved_filename and file_path:
            resolved_filename = os.path.basename(file_path)

        resolved_bytes = file_bytes
        resolved_content = content or (content_or_bytes if isinstance(content_or_bytes, str) else None)

        if resolved_bytes is None and isinstance(content_or_bytes, bytes):
            resolved_bytes = content_or_bytes

        # Read bytes from file path if available
        if resolved_bytes is None and file_path and os.path.exists(file_path) and os.path.isfile(file_path):
            try:
                with open(file_path, "rb") as f:
                    resolved_bytes = f.read()
            except Exception as read_err:
                logger.warning(f"[DocumentRouter] Could not read file from path {file_path}: {read_err}")

        # Check for base64 or path in string content
        ext = os.path.splitext(resolved_filename)[1].lower() if resolved_filename else ""

        if resolved_bytes is None and resolved_content:
            if os.path.exists(resolved_content) and os.path.isfile(resolved_content):
                try:
                    with open(resolved_content, "rb") as f:
                        resolved_bytes = f.read()
                    if not resolved_filename:
                        resolved_filename = os.path.basename(resolved_content)
                        ext = os.path.splitext(resolved_filename)[1].lower()
                except Exception:
                    pass
            elif ext in (".pdf", ".docx", ".pptx"):
                import base64

                clean_str = resolved_content.strip()
                if clean_str.startswith("data:") and ";base64," in clean_str:
                    clean_str = clean_str.split(";base64,")[1]
                try:
                    decoded = base64.b64decode(clean_str, validate=True)
                    if len(decoded) > 0:
                        resolved_bytes = decoded
                except Exception:
                    resolved_bytes = resolved_content.encode("latin1", errors="replace")

        # Route by file extension
        if ext == ".pptx" and resolved_bytes:
            return self._process_pptx(resolved_bytes, resolved_filename)

        if ext == ".docx" and resolved_bytes:
            return self._process_docx(resolved_bytes, resolved_filename)

        if ext == ".pdf" and resolved_bytes:
            return self._process_pdf(resolved_bytes, resolved_filename)

        # Fallback to local processor for markdown, text, or non-binary content
        text_to_process = resolved_content
        if text_to_process is None and resolved_bytes:
            text_to_process = resolved_bytes.decode("utf-8", errors="replace")

        clean_md = self.fallback_processor.convert_to_markdown(
            content_or_bytes=text_to_process,
            file_path=file_path,
            file_bytes=resolved_bytes,
            filename=resolved_filename,
        )
        return self.fallback_processor.process_markdown(clean_md)

    def _process_pptx(
        self,
        file_bytes: bytes,
        filename: str,
    ) -> Tuple[List[Dict[str, Any]], Dict[str, List[str]]]:
        """Extract slide-by-slide units and speaker notes from PPTX presentation."""
        try:
            import pptx
            stream = io.BytesIO(file_bytes)
            prs = pptx.Presentation(stream)
        except Exception as pptx_err:
            logger.warning(
                f"[DocumentRouter] python-pptx failed to load presentation: {pptx_err}. Falling back."
            )
            clean_md = self.fallback_processor.convert_to_markdown(file_bytes=file_bytes, filename=filename)
            return self.fallback_processor.process_markdown(clean_md)

        final_documents: List[Dict[str, Any]] = []
        toc: Dict[str, List[str]] = {}

        topic = os.path.splitext(filename)[0] if filename else "Slide Presentation"
        if len(prs.slides) > 0:
            first_slide = prs.slides[0]
            if first_slide.shapes.title and first_slide.shapes.title.text.strip():
                topic = first_slide.shapes.title.text.strip()
            else:
                for shp in first_slide.shapes:
                    if hasattr(shp, "text") and shp.text.strip():
                        first_line = shp.text.strip().splitlines()[0].strip()
                        if first_line:
                            topic = first_line
                            break

        toc[topic] = []
        global_seq_id = 0

        for idx, slide in enumerate(prs.slides, 1):
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text.strip():
                slide_title = slide.shapes.title.text.strip()

            body_lines: List[str] = []
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if hasattr(shape, "text") and shape.text.strip():
                    body_lines.append(shape.text.strip())

            if not slide_title:
                if body_lines:
                    first_line = body_lines[0].splitlines()[0].strip()
                    if len(first_line) < 60:
                        slide_title = first_line
                        body_lines = body_lines[1:]
                if not slide_title:
                    slide_title = f"Slide {idx}"

            # Extract speaker notes
            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide:
                notes_tf = slide.notes_slide.notes_text_frame
                if notes_tf and notes_tf.text.strip():
                    notes_text = notes_tf.text.strip()

            slide_body = "\n\n".join(body_lines).strip()
            if notes_text:
                if slide_body:
                    slide_body = f"{slide_body}\n\n**Speaker Notes:**\n{notes_text}"
                else:
                    slide_body = f"**Speaker Notes:**\n{notes_text}"

            if not slide_body:
                slide_body = slide_title

            if slide_title not in toc[topic]:
                toc[topic].append(slide_title)

            # Slicing: Exactly 1 chunk per slide, sub-chunk only if > 1200 words
            sub_chunks = sub_chunk_section(
                section_text=slide_body,
                chapter=topic,
                section=slide_title,
                max_words=1200,
            )

            for sub_idx, sub_text in enumerate(sub_chunks):
                final_documents.append({
                    "page_content": sub_text.strip(),
                    "metadata": {
                        "Chapter": topic,
                        "Section": slide_title,
                        "seq_id": global_seq_id,
                        "document_type": "SLIDE",
                        "slide_number": idx,
                        "sub_chunk_index": sub_idx,
                        "total_sub_chunks": len(sub_chunks),
                    },
                })
                global_seq_id += 1

        if not final_documents:
            clean_md = self.fallback_processor.convert_to_markdown(file_bytes=file_bytes, filename=filename)
            return self.fallback_processor.process_markdown(clean_md)

        return final_documents, toc

    def _process_docx(
        self,
        file_bytes: bytes,
        filename: str,
    ) -> Tuple[List[Dict[str, Any]], Dict[str, List[str]]]:
        """Extract native heading hierarchy (Heading 1/2/3) from DOCX document."""
        try:
            import docx
            stream = io.BytesIO(file_bytes)
            doc = docx.Document(stream)
        except Exception as docx_err:
            logger.warning(
                f"[DocumentRouter] python-docx failed to load document: {docx_err}. Falling back."
            )
            clean_md = self.fallback_processor.convert_to_markdown(file_bytes=file_bytes, filename=filename)
            return self.fallback_processor.process_markdown(clean_md)

        final_documents: List[Dict[str, Any]] = []
        toc: Dict[str, List[str]] = {}

        doc_title = os.path.splitext(filename)[0] if filename else "Document"
        current_chapter = doc_title
        current_section = "Overview"
        current_subsection = ""

        nodes_content = []
        current_blocks: List[str] = []

        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue
            style_name = (p.style.name if p.style else "").lower()
            if "heading 1" in style_name:
                if current_blocks:
                    nodes_content.append((current_chapter, current_section, current_subsection, list(current_blocks)))
                    current_blocks = []
                current_chapter = txt
                current_section = f"{txt} (Overview)"
                current_subsection = ""
            elif "heading 2" in style_name:
                if current_blocks:
                    nodes_content.append((current_chapter, current_section, current_subsection, list(current_blocks)))
                    current_blocks = []
                current_section = txt
                current_subsection = ""
            elif "heading 3" in style_name:
                if current_blocks:
                    nodes_content.append((current_chapter, current_section, current_subsection, list(current_blocks)))
                    current_blocks = []
                current_subsection = txt
            else:
                current_blocks.append(txt)

        if current_blocks:
            nodes_content.append((current_chapter, current_section, current_subsection, list(current_blocks)))

        global_seq_id = 0
        for ch, sec, sub, blocks in nodes_content:
            content_text = "\n\n".join(blocks).strip()
            if not content_text:
                continue

            if ch not in toc:
                toc[ch] = []

            display_sec = f"{sec} > {sub}" if sub else sec
            if display_sec not in toc[ch]:
                toc[ch].append(display_sec)

            sub_chunks = sub_chunk_section(
                section_text=content_text,
                chapter=ch,
                section=display_sec,
                max_words=1200,
            )

            for sub_idx, sub_text in enumerate(sub_chunks):
                final_documents.append({
                    "page_content": sub_text.strip(),
                    "metadata": {
                        "Chapter": ch,
                        "Section": display_sec,
                        "seq_id": global_seq_id,
                        "document_type": "BOOK" if sub else "PAPER",
                        "sub_chunk_index": sub_idx,
                        "total_sub_chunks": len(sub_chunks),
                    },
                })
                global_seq_id += 1

        if not final_documents:
            clean_md = self.fallback_processor.convert_to_markdown(file_bytes=file_bytes, filename=filename)
            return self.fallback_processor.process_markdown(clean_md)

        return final_documents, toc

    def _process_pdf(
        self,
        file_bytes: bytes,
        filename: str,
    ) -> Tuple[List[Dict[str, Any]], Dict[str, List[str]]]:
        """Process PDF document via Gemini Multimodal Vision with deterministic page bounds validation."""
        total_pages = 1
        pdf_doc = None
        try:
            import fitz
            pdf_doc = fitz.open(stream=file_bytes, filetype="pdf")
            total_pages = pdf_doc.page_count
        except Exception as fitz_err:
            logger.warning(f"[DocumentRouter] fitz could not open PDF stream: {fitz_err}")

        # Attempt Gemini Multimodal Vision extraction
        structure: Optional[DocumentStructure] = None
        if self.gemini_extractor:
            structure = self.gemini_extractor.extract_structure(
                file_bytes=file_bytes,
                total_pages=total_pages,
            )

        # Validate and sanitize page boundaries
        if structure and self.page_validator:
            structure = self.page_validator.validate_and_sanitize(
                structure=structure,
                total_pages=total_pages,
            )

        # Slice PDF according to validated DocumentStructure
        if structure and structure.hierarchy and pdf_doc is not None:
            try:
                final_documents, toc = self._slice_pdf_hierarchy(pdf_doc, structure)
                if final_documents:
                    logger.info(
                        f"[DocumentRouter] Successfully routed PDF as {structure.document_type.value} with {len(final_documents)} chunks."
                    )
                    return final_documents, toc
            except Exception as slice_err:
                logger.warning(
                    f"[DocumentRouter] Slicing error: {slice_err}. Falling back to local processor."
                )

        # Fallback to local typographic processor
        logger.info("[DocumentRouter] Falling back to MathAwareDocumentProcessor for PDF.")
        clean_md = self.fallback_processor.convert_to_markdown(file_bytes=file_bytes, filename=filename)
        return self.fallback_processor.process_markdown(clean_md)

    def _slice_pdf_hierarchy(
        self,
        pdf_doc: Any,
        structure: DocumentStructure,
    ) -> Tuple[List[Dict[str, Any]], Dict[str, List[str]]]:
        """Slice PDF content into pedagogical chunks based on document type and hierarchy."""
        final_documents: List[Dict[str, Any]] = []
        toc: Dict[str, List[str]] = {}
        global_seq_id = 0

        doc_type = structure.document_type
        doc_title = structure.document_title or "Document"

        def get_pages_text(start_p: int, end_p: int) -> str:
            texts = []
            s = max(1, min(start_p, pdf_doc.page_count))
            e = max(s, min(end_p, pdf_doc.page_count))
            for p_num in range(s, e + 1):
                try:
                    p = pdf_doc[p_num - 1]
                    t = p.get_text()
                    if t and t.strip():
                        texts.append(t.strip())
                except Exception:
                    pass
            return "\n\n".join(texts)

        def partition_page_text_by_titles(
            page_num: int,
            nodes: List[Any],
        ) -> Dict[str, str]:
            """Partition text of a single physical page among candidate nodes based on their titles.

            Preserves preamble text before the first section heading in the first chunk,
            and handles partial title matches without discarding unmatched sections.
            """
            p_idx = max(1, min(page_num, pdf_doc.page_count)) - 1
            try:
                page_text = pdf_doc[p_idx].get_text()
            except Exception:
                return {}

            if not page_text or not page_text.strip():
                return {}

            locations = []
            lower_page = page_text.lower()
            for node in nodes:
                node_title = getattr(node, "title", None) or getattr(node, "name", "")
                title = node_title.strip()
                if not title:
                    continue
                # Try exact case-insensitive match
                idx = lower_page.find(title.lower())
                if idx != -1:
                    locations.append((idx, title, node))
                else:
                    # Fallback to word prefix match (4, 3, 2, or 1 words)
                    words = title.split()
                    for count in (4, 3, 2, 1):
                        if len(words) >= count:
                            short_title = " ".join(words[:count]).lower()
                            idx = lower_page.find(short_title)
                            if idx != -1:
                                locations.append((idx, title, node))
                                break

            # Deduplicate by title, preserving earliest occurrence
            unique_locations = []
            seen_titles = set()
            locations.sort(key=lambda x: x[0])
            for idx, title, node in locations:
                if title not in seen_titles:
                    seen_titles.add(title)
                    unique_locations.append((idx, title, node))

            result = {}
            if unique_locations:
                for i, (idx, title, node) in enumerate(unique_locations):
                    # CRITICAL FIX: start_idx is 0 for first section on page to preserve preamble
                    start_idx = 0 if i == 0 else idx
                    end_idx = unique_locations[i + 1][0] if i + 1 < len(unique_locations) else len(page_text)
                    chunk_text = page_text[start_idx:end_idx].strip()
                    if chunk_text:
                        result[title] = chunk_text

            return result

        if doc_type == DocumentType.SLIDE:
            # SLIDE: Topic > Slide Title (1 atomic chunk per slide)
            for chapter in structure.hierarchy:
                topic = chapter.name or doc_title
                if topic not in toc:
                    toc[topic] = []

                for section in chapter.sections:
                    slide_title = section.title
                    slide_text = get_pages_text(section.start_page, section.end_page)
                    if not slide_text.strip():
                        slide_text = f"{slide_title}\n[Slide {section.start_page}]"

                    if slide_title not in toc[topic]:
                        toc[topic].append(slide_title)

                    sub_chunks = sub_chunk_section(
                        section_text=slide_text,
                        chapter=topic,
                        section=slide_title,
                        max_words=1200,
                    )

                    for sub_idx, sub_text in enumerate(sub_chunks):
                        final_documents.append({
                            "page_content": sub_text.strip(),
                            "metadata": {
                                "Chapter": topic,
                                "Section": slide_title,
                                "seq_id": global_seq_id,
                                "document_type": "SLIDE",
                                "start_page": section.start_page,
                                "end_page": section.end_page,
                                "sub_chunk_index": sub_idx,
                                "total_sub_chunks": len(sub_chunks),
                            },
                        })
                        global_seq_id += 1

        elif doc_type == DocumentType.BOOK:
            # BOOK: Deep Chapter > Section > Subsection slicing
            for chapter in structure.hierarchy:
                chap_name = chapter.name
                if chap_name not in toc:
                    toc[chap_name] = []

                for section in chapter.sections:
                    sec_name = section.title
                    subsections = section.subsections or []

                    if subsections:
                        # CRITICAL FIX: Group same-page subsections by page number to support
                        # multiple distinct single-page subsection groups in the same section.
                        page_to_subs: Dict[int, List[SubSectionNode]] = {}
                        for s in subsections:
                            if s.start_page == s.end_page:
                                page_to_subs.setdefault(s.start_page, []).append(s)

                        same_page_map: Dict[str, str] = {}
                        for p_num, p_subs in page_to_subs.items():
                            if len(p_subs) > 1 or any(sb.is_same_page for sb in p_subs):
                                partitioned = partition_page_text_by_titles(p_num, p_subs)
                                same_page_map.update(partitioned)

                        for sub in subsections:
                            sub_title = sub.title
                            full_sec_path = f"{sec_name} > {sub_title}"

                            if sub_title in same_page_map and same_page_map[sub_title]:
                                sub_text = same_page_map[sub_title]
                            else:
                                sub_text = get_pages_text(sub.start_page, sub.end_page)

                            if not sub_text.strip():
                                sub_text = f"{full_sec_path}\n[Pages {sub.start_page}-{sub.end_page}]"

                            if full_sec_path not in toc[chap_name]:
                                toc[chap_name].append(full_sec_path)

                            # Context: [Pedagogical Context: {Chapter} > {Section} > {Subsection}]
                            sub_chunks = sub_chunk_section(
                                section_text=sub_text,
                                chapter=chap_name,
                                section=full_sec_path,
                                max_words=1200,
                            )

                            for sub_idx, sub_chunk in enumerate(sub_chunks):
                                final_documents.append({
                                    "page_content": sub_chunk.strip(),
                                    "metadata": {
                                        "Chapter": chap_name,
                                        "Section": full_sec_path,
                                        "seq_id": global_seq_id,
                                        "document_type": "BOOK",
                                        "start_page": sub.start_page,
                                        "end_page": sub.end_page,
                                        "sub_chunk_index": sub_idx,
                                        "total_sub_chunks": len(sub_chunks),
                                    },
                                })
                                global_seq_id += 1
                    else:
                        sec_text = get_pages_text(section.start_page, section.end_page)
                        if not sec_text.strip():
                            sec_text = f"{sec_name}\n[Pages {section.start_page}-{section.end_page}]"

                        if sec_name not in toc[chap_name]:
                            toc[chap_name].append(sec_name)

                        sub_chunks = sub_chunk_section(
                            section_text=sec_text,
                            chapter=chap_name,
                            section=sec_name,
                            max_words=1200,
                        )

                        for sub_idx, sub_chunk in enumerate(sub_chunks):
                            final_documents.append({
                                "page_content": sub_chunk.strip(),
                                "metadata": {
                                    "Chapter": chap_name,
                                    "Section": sec_name,
                                    "seq_id": global_seq_id,
                                    "document_type": "BOOK",
                                    "start_page": section.start_page,
                                    "end_page": section.end_page,
                                    "sub_chunk_index": sub_idx,
                                    "total_sub_chunks": len(sub_chunks),
                                },
                            })
                            global_seq_id += 1

        else:
            # PAPER: [Pedagogical Context: {Paper Title} > {Section}]
            paper_title = doc_title

            # Pre-partition root chapters sharing a single physical page (when chapters lack child sections)
            chap_page_map: Dict[str, str] = {}
            chap_by_page: Dict[int, List[Any]] = {}
            for chap in structure.hierarchy:
                if not chap.sections and chap.start_page == chap.end_page:
                    chap_by_page.setdefault(chap.start_page, []).append(chap)
            for p_num, c_nodes in chap_by_page.items():
                if len(c_nodes) > 1:
                    chap_page_map.update(partition_page_text_by_titles(p_num, c_nodes))

            for chapter in structure.hierarchy:
                major_sec = chapter.name
                if not chapter.sections:
                    if major_sec in chap_page_map and chap_page_map[major_sec]:
                        sec_text = chap_page_map[major_sec]
                    else:
                        sec_text = get_pages_text(chapter.start_page, chapter.end_page)

                    if not sec_text.strip():
                        sec_text = f"{major_sec}\n[Pages {chapter.start_page}-{chapter.end_page}]"

                    if paper_title not in toc:
                        toc[paper_title] = []
                    if major_sec not in toc[paper_title]:
                        toc[paper_title].append(major_sec)

                    sub_chunks = sub_chunk_section(
                        section_text=sec_text,
                        chapter=paper_title,
                        section=major_sec,
                        max_words=1200,
                    )

                    for sub_idx, sub_chunk in enumerate(sub_chunks):
                        final_documents.append({
                            "page_content": sub_chunk.strip(),
                            "metadata": {
                                "Chapter": paper_title,
                                "Section": major_sec,
                                "seq_id": global_seq_id,
                                "document_type": "PAPER",
                                "start_page": chapter.start_page,
                                "end_page": chapter.end_page,
                                "sub_chunk_index": sub_idx,
                                "total_sub_chunks": len(sub_chunks),
                            },
                        })
                        global_seq_id += 1
                else:
                    if major_sec not in toc:
                        toc[major_sec] = []

                    # Pre-partition sibling sections sharing the same physical page
                    sec_page_map: Dict[str, str] = {}
                    sec_by_page: Dict[int, List[Any]] = {}
                    for s in chapter.sections:
                        if not s.subsections and s.start_page == s.end_page:
                            sec_by_page.setdefault(s.start_page, []).append(s)
                    for p_num, s_nodes in sec_by_page.items():
                        if len(s_nodes) > 1 or any(sn.is_same_page for sn in s_nodes):
                            sec_page_map.update(partition_page_text_by_titles(p_num, s_nodes))

                    for section in chapter.sections:
                        sec_name = section.title
                        subsections = section.subsections or []

                        if subsections:
                            # Pre-partition sibling subsections sharing the same physical page
                            sub_page_map: Dict[str, str] = {}
                            sub_by_page: Dict[int, List[SubSectionNode]] = {}
                            for sub in subsections:
                                if sub.start_page == sub.end_page:
                                    sub_by_page.setdefault(sub.start_page, []).append(sub)
                            for p_num, sub_nodes in sub_by_page.items():
                                if len(sub_nodes) > 1 or any(sb.is_same_page for sb in sub_nodes):
                                    sub_page_map.update(partition_page_text_by_titles(p_num, sub_nodes))

                            for sub in subsections:
                                sub_name = sub.title
                                full_sec = f"{sec_name} > {sub_name}"

                                if sub_name in sub_page_map and sub_page_map[sub_name]:
                                    sub_text = sub_page_map[sub_name]
                                else:
                                    sub_text = get_pages_text(sub.start_page, sub.end_page)

                                if not sub_text.strip():
                                    sub_text = f"{full_sec}\n[Pages {sub.start_page}-{sub.end_page}]"

                                if full_sec not in toc[major_sec]:
                                    toc[major_sec].append(full_sec)

                                sub_chunks = sub_chunk_section(
                                    section_text=sub_text,
                                    chapter=paper_title,
                                    section=f"{major_sec} > {full_sec}",
                                    max_words=1200,
                                )

                                for sub_idx, sub_chunk in enumerate(sub_chunks):
                                    final_documents.append({
                                        "page_content": sub_chunk.strip(),
                                        "metadata": {
                                            "Chapter": paper_title,
                                            "Section": f"{major_sec} > {full_sec}",
                                            "seq_id": global_seq_id,
                                            "document_type": "PAPER",
                                            "start_page": sub.start_page,
                                            "end_page": sub.end_page,
                                            "sub_chunk_index": sub_idx,
                                            "total_sub_chunks": len(sub_chunks),
                                        },
                                    })
                                    global_seq_id += 1
                        else:
                            if sec_name in sec_page_map and sec_page_map[sec_name]:
                                sec_text = sec_page_map[sec_name]
                            else:
                                sec_text = get_pages_text(section.start_page, section.end_page)

                            if not sec_text.strip():
                                sec_text = f"{sec_name}\n[Pages {section.start_page}-{section.end_page}]"

                            if sec_name not in toc[major_sec]:
                                toc[major_sec].append(sec_name)

                            sub_chunks = sub_chunk_section(
                                section_text=sec_text,
                                chapter=paper_title,
                                section=f"{major_sec} > {sec_name}",
                                max_words=1200,
                            )

                            for sub_idx, sub_chunk in enumerate(sub_chunks):
                                final_documents.append({
                                    "page_content": sub_chunk.strip(),
                                    "metadata": {
                                        "Chapter": paper_title,
                                        "Section": f"{major_sec} > {sec_name}",
                                        "seq_id": global_seq_id,
                                        "document_type": "PAPER",
                                        "start_page": section.start_page,
                                        "end_page": section.end_page,
                                        "sub_chunk_index": sub_idx,
                                        "total_sub_chunks": len(sub_chunks),
                                    },
                                })
                                global_seq_id += 1

        return final_documents, toc

    def convert_to_markdown(
        self,
        content_or_bytes: Union[str, bytes, None] = None,
        file_path: Optional[str] = None,
        file_bytes: Optional[bytes] = None,
        filename: str = "",
    ) -> str:
        """Compatibility delegate to MathAwareDocumentProcessor.convert_to_markdown."""
        return self.fallback_processor.convert_to_markdown(
            content_or_bytes=content_or_bytes,
            file_path=file_path,
            file_bytes=file_bytes,
            filename=filename,
        )

    def process_markdown(
        self,
        markdown_text: str,
        filename: str = "",
    ) -> Tuple[List[Dict[str, Any]], Dict[str, List[str]]]:
        """Compatibility delegate to MathAwareDocumentProcessor.process_markdown."""
        return self.fallback_processor.process_markdown(markdown_text)

    def process_file(
        self,
        file_path: str,
    ) -> Tuple[List[Dict[str, Any]], Dict[str, List[str]]]:
        """Process document file on disk."""
        filename = os.path.basename(file_path)
        with open(file_path, "rb") as f:
            file_bytes = f.read()
        return self.process_document(
            file_path=file_path,
            file_bytes=file_bytes,
            filename=filename,
        )
