"""Unit tests for DocumentRouter across PPTX, DOCX, and PDF with Gemini multimodal vision."""

import io
from unittest.mock import MagicMock
import fitz
import docx
import pptx
import pytest

from core.schemas import (
    ChapterOrTopicNode,
    DocumentStructure,
    DocumentType,
    SectionNode,
    SubSectionNode,
)
from database.document_router import DocumentRouter
from database.gemini_extractor import GeminiStructureExtractor
from database.page_validator import PageBoundsValidator


@pytest.fixture
def mock_gemini():
    extractor = MagicMock(spec=GeminiStructureExtractor)
    extractor.extract_structure.return_value = None
    return extractor


@pytest.fixture
def page_validator():
    return PageBoundsValidator()


@pytest.fixture
def router(mock_gemini, page_validator):
    return DocumentRouter(
        gemini_extractor=mock_gemini,
        page_validator=page_validator,
    )


def test_pptx_routing_slide_by_slide_and_notes(router):
    """PPTX files should be parsed slide-by-slide with speaker notes preserved."""
    prs = pptx.Presentation()
    blank_layout = prs.slide_layouts[6]  # Blank layout

    # Slide 1: Title and body
    slide1 = prs.slides.add_slide(blank_layout)
    tx_box1 = slide1.shapes.add_textbox(10, 10, 500, 100)
    tx_box1.text_frame.text = "Introduction to Deep Learning\nFoundations and History"

    # Slide 2: Title, body, and speaker notes
    slide2 = prs.slides.add_slide(blank_layout)
    tx_box2 = slide2.shapes.add_textbox(10, 10, 500, 100)
    tx_box2.text_frame.text = "Backpropagation Algorithm\nDerivation of gradients"
    notes_slide = slide2.notes_slide
    notes_slide.notes_text_frame.text = "Explain the chain rule carefully."

    pptx_bytes_io = io.BytesIO()
    prs.save(pptx_bytes_io)
    pptx_bytes = pptx_bytes_io.getvalue()

    chunks, toc = router.process_document(
        file_bytes=pptx_bytes,
        filename="lecture_01.pptx",
    )

    assert len(chunks) == 2
    assert "Introduction to Deep Learning" in toc
    assert len(toc["Introduction to Deep Learning"]) == 2

    # Verify speaker notes are preserved in slide 2
    slide2_chunk = chunks[1]
    assert "Explain the chain rule carefully." in slide2_chunk["page_content"]
    assert "Backpropagation Algorithm" in slide2_chunk["metadata"]["Section"]
    assert slide2_chunk["metadata"]["document_type"] == "SLIDE"
    assert "[Pedagogical Context:" in slide2_chunk["page_content"]


def test_docx_routing_native_headings(router):
    """DOCX files should be parsed using native heading hierarchy."""
    doc = docx.Document()
    doc.add_heading("Optimization Techniques", level=1)
    doc.add_paragraph("Overview of first-order optimization.")
    doc.add_heading("Gradient Descent", level=2)
    doc.add_paragraph("Batch gradient descent computes the gradient over the entire dataset.")
    doc.add_heading("Learning Rate Schedule", level=3)
    doc.add_paragraph("Step decay and cosine annealing are popular schedules.")

    docx_bytes_io = io.BytesIO()
    doc.save(docx_bytes_io)
    docx_bytes = docx_bytes_io.getvalue()

    chunks, toc = router.process_document(
        file_bytes=docx_bytes,
        filename="optimization.docx",
    )

    assert len(chunks) >= 2
    assert "Optimization Techniques" in toc
    # Check pedagogical context
    found_context = any("[Pedagogical Context: Optimization Techniques >" in c["page_content"] for c in chunks)
    assert found_context


def test_pdf_routing_slide_type(mock_gemini, page_validator):
    """PDF classified as SLIDE should slice into atomic slides with [Context: Topic > Slide]."""
    # Create a 2-page in-memory PDF
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "Slide 1: Transformers Architecture Overview")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "Slide 2: Multi-Head Self-Attention Mechanism")
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    mock_gemini.extract_structure.return_value = DocumentStructure(
        document_type=DocumentType.SLIDE,
        document_title="Attention Is All You Need Presentation",
        total_pages=2,
        summary="Overview of transformer architecture.",
        hierarchy=[
            ChapterOrTopicNode(
                name="Module 1: Attention",
                start_page=1,
                end_page=2,
                sections=[
                    SectionNode(title="Slide 1: Architecture", start_page=1, end_page=1),
                    SectionNode(title="Slide 2: Self-Attention", start_page=2, end_page=2),
                ],
            )
        ],
    )

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="transformers_deck.pdf")

    assert len(chunks) == 2
    assert chunks[0]["metadata"]["document_type"] == "SLIDE"
    assert "[Pedagogical Context: Module 1: Attention > Slide 1: Architecture]" in chunks[0]["page_content"]
    assert "[Pedagogical Context: Module 1: Attention > Slide 2: Self-Attention]" in chunks[1]["page_content"]
    assert "Module 1: Attention" in toc


def test_pdf_routing_paper_type(mock_gemini, page_validator):
    """PDF classified as PAPER should slice by sections with [Context: Paper Title > Section]."""
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "1. Introduction\nRecent advances in LLM reasoning.")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "2. Methodology\nWe evaluate chain-of-thought prompting.")
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    mock_gemini.extract_structure.return_value = DocumentStructure(
        document_type=DocumentType.PAPER,
        document_title="Empirical Reasoning in LLMs",
        total_pages=2,
        summary="A study on prompting methods.",
        hierarchy=[
            ChapterOrTopicNode(
                name="1. Introduction",
                start_page=1,
                end_page=1,
                sections=[],
            ),
            ChapterOrTopicNode(
                name="2. Methodology",
                start_page=2,
                end_page=2,
                sections=[],
            ),
        ],
    )

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="paper.pdf")

    assert len(chunks) == 2
    assert chunks[0]["metadata"]["document_type"] == "PAPER"
    assert "[Pedagogical Context: Empirical Reasoning in LLMs > 1. Introduction]" in chunks[0]["page_content"]
    assert "[Pedagogical Context: Empirical Reasoning in LLMs > 2. Methodology]" in chunks[1]["page_content"]


def test_pdf_routing_book_deep_subsection_chunking(mock_gemini, page_validator):
    """PDF classified as BOOK should slice deeply down to Chapter > Section > Subsection."""
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "1.1.1 Stochastic Gradient Descent\nSGD updates weights per mini-batch.")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "1.1.2 Momentum Optimization\nMomentum accelerates in consistent directions.")
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    mock_gemini.extract_structure.return_value = DocumentStructure(
        document_type=DocumentType.BOOK,
        document_title="Deep Learning Textbook",
        total_pages=2,
        summary="Comprehensive textbook on deep learning.",
        hierarchy=[
            ChapterOrTopicNode(
                name="Chapter 1: Optimization",
                start_page=1,
                end_page=2,
                sections=[
                    SectionNode(
                        title="1.1 First-Order Methods",
                        start_page=1,
                        end_page=2,
                        subsections=[
                            SubSectionNode(
                                title="1.1.1 Stochastic Gradient Descent",
                                start_page=1,
                                end_page=1,
                            ),
                            SubSectionNode(
                                title="1.1.2 Momentum Optimization",
                                start_page=2,
                                end_page=2,
                            ),
                        ],
                    )
                ],
            )
        ],
    )

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="deep_learning_book.pdf")

    assert len(chunks) == 2
    assert chunks[0]["metadata"]["document_type"] == "BOOK"
    assert "[Pedagogical Context: Chapter 1: Optimization > 1.1 First-Order Methods > 1.1.1 Stochastic Gradient Descent]" in chunks[0]["page_content"]
    assert "[Pedagogical Context: Chapter 1: Optimization > 1.1 First-Order Methods > 1.1.2 Momentum Optimization]" in chunks[1]["page_content"]
    assert "Chapter 1: Optimization" in toc


def test_pdf_same_page_subsection_partitioning(mock_gemini, page_validator):
    """Subsections located on the exact same page should be partitioned by title."""
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text(
        (50, 50),
        "1.1 QuickSort\nQuickSort partitions around a pivot.\n\n"
        "1.2 MergeSort\nMergeSort divides the array in half and merges sorted sub-arrays."
    )
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    mock_gemini.extract_structure.return_value = DocumentStructure(
        document_type=DocumentType.BOOK,
        document_title="Algorithms Guide",
        total_pages=1,
        summary="Guide to algorithms.",
        hierarchy=[
            ChapterOrTopicNode(
                name="Chapter 1: Divide and Conquer",
                start_page=1,
                end_page=1,
                sections=[
                    SectionNode(
                        title="Sorting",
                        start_page=1,
                        end_page=1,
                        subsections=[
                            SubSectionNode(title="1.1 QuickSort", start_page=1, end_page=1),
                            SubSectionNode(title="1.2 MergeSort", start_page=1, end_page=1),
                        ],
                    )
                ],
            )
        ],
    )

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="algorithms.pdf")

    assert len(chunks) == 2
    assert "QuickSort partitions around a pivot" in chunks[0]["page_content"]
    assert "MergeSort divides the array" in chunks[1]["page_content"]


def test_pdf_fallback_to_local_processor_when_gemini_fails(mock_gemini, page_validator):
    """When Gemini returns None, router should fall back gracefully to local processor."""
    mock_gemini.extract_structure.return_value = None

    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "1. Introduction\nFallback test content.")
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="fallback_doc.pdf")

    assert len(chunks) >= 1
    assert len(toc) >= 1


def test_sub_chunk_section_long_text_split(router):
    """Sections exceeding 1,200 words must be sub-chunked with context headers."""
    long_content = "Word " * 1500
    chunks = router.fallback_processor.sub_chunk_section(
        section_text=long_content,
        chapter="Physics",
        section="Mechanics",
        max_words=1200,
        target_chunk_words=700,
        overlap_words=150,
    )
    assert len(chunks) >= 2
    for chunk in chunks:
        assert chunk.startswith("[Pedagogical Context: Physics > Mechanics]")


def test_pdf_same_page_preamble_preservation(mock_gemini, page_validator):
    """Preamble text before the first subsection heading must not be lost."""
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text(
        (50, 50),
        "Important Introductory Overview of Sorting\n\n"
        "1.1 QuickSort\nQuickSort partitions around a pivot.\n\n"
        "1.2 MergeSort\nMergeSort divides in half."
    )
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    mock_gemini.extract_structure.return_value = DocumentStructure(
        document_type=DocumentType.BOOK,
        document_title="Algorithms Guide",
        total_pages=1,
        hierarchy=[
            ChapterOrTopicNode(
                name="Chapter 1: Sorting",
                start_page=1,
                end_page=1,
                sections=[
                    SectionNode(
                        title="Algorithms",
                        start_page=1,
                        end_page=1,
                        subsections=[
                            SubSectionNode(title="1.1 QuickSort", start_page=1, end_page=1),
                            SubSectionNode(title="1.2 MergeSort", start_page=1, end_page=1),
                        ],
                    )
                ],
            )
        ],
    )

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="algo.pdf")

    assert len(chunks) == 2
    # Verify the introductory preamble is preserved in the first subsection
    assert "Important Introductory Overview" in chunks[0]["page_content"]
    assert "QuickSort partitions around a pivot" in chunks[0]["page_content"]
    assert "MergeSort divides in half" in chunks[1]["page_content"]


def test_pdf_multi_page_same_page_subsections(mock_gemini, page_validator):
    """Subsections grouped on distinct physical pages (e.g. 2 on page 1, 2 on page 2) must all partition."""
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "1.1 A\nAlpha text\n\n1.2 B\nBeta text")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "1.3 C\nGamma text\n\n1.4 D\nDelta text")
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    mock_gemini.extract_structure.return_value = DocumentStructure(
        document_type=DocumentType.BOOK,
        document_title="Multi Page Textbook",
        total_pages=2,
        hierarchy=[
            ChapterOrTopicNode(
                name="Chapter 1",
                start_page=1,
                end_page=2,
                sections=[
                    SectionNode(
                        title="Section 1",
                        start_page=1,
                        end_page=2,
                        subsections=[
                            SubSectionNode(title="1.1 A", start_page=1, end_page=1),
                            SubSectionNode(title="1.2 B", start_page=1, end_page=1),
                            SubSectionNode(title="1.3 C", start_page=2, end_page=2),
                            SubSectionNode(title="1.4 D", start_page=2, end_page=2),
                        ],
                    )
                ],
            )
        ],
    )

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="multi_page.pdf")

    assert len(chunks) == 4
    # Ensure chunk 2 does not have duplicate chunk 1 text, and chunk 4 does not have duplicate chunk 3 text
    assert "Alpha text" in chunks[0]["page_content"]
    assert "Beta text" in chunks[1]["page_content"]
    assert "Alpha text" not in chunks[1]["page_content"]
    assert "Gamma text" in chunks[2]["page_content"]
    assert "Delta text" in chunks[3]["page_content"]
    assert "Gamma text" not in chunks[3]["page_content"]


def test_pdf_paper_same_page_sections_partitioned(mock_gemini, page_validator):
    """2-column academic PAPER with multiple sections on page 1 must partition cleanly."""
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text(
        (50, 50),
        "1. Introduction\nIntroductory background on deep learning.\n\n"
        "2. Related Work\nPrior studies on attention architectures."
    )
    pdf_bytes = pdf_doc.tobytes()
    pdf_doc.close()

    mock_gemini.extract_structure.return_value = DocumentStructure(
        document_type=DocumentType.PAPER,
        document_title="Attention Paper",
        total_pages=1,
        hierarchy=[
            ChapterOrTopicNode(
                name="1. Introduction",
                start_page=1,
                end_page=1,
                sections=[],
            ),
            ChapterOrTopicNode(
                name="2. Related Work",
                start_page=1,
                end_page=1,
                sections=[],
            ),
        ],
    )

    router = DocumentRouter(gemini_extractor=mock_gemini, page_validator=page_validator)
    chunks, toc = router.process_document(file_bytes=pdf_bytes, filename="paper_col.pdf")

    assert len(chunks) == 2
    assert "Introductory background" in chunks[0]["page_content"]
    assert "Prior studies on attention" in chunks[1]["page_content"]
    assert "Prior studies on attention" not in chunks[0]["page_content"]


def test_gemini_extractor_parses_dict_response():
    """GeminiStructureExtractor should parse response.parsed when returned as a dict."""
    extractor = GeminiStructureExtractor(api_key="mock_key")
    mock_client = MagicMock()
    mock_resp = MagicMock()
    mock_resp.parsed = {
        "document_type": "PAPER",
        "document_title": "Parsed Dict Paper",
        "total_pages": 3,
        "summary": "Pedagogical summary from dict.",
        "hierarchy": [
            {
                "name": "1. Intro",
                "start_page": 1,
                "end_page": 2,
                "sections": [],
            }
        ],
    }
    mock_resp.text = ""
    mock_client.models.generate_content.return_value = mock_resp
    extractor._client = mock_client

    result = extractor.extract_structure(file_bytes=b"%PDF-1.4 dummy", total_pages=3)
    assert result is not None
    assert result.document_type == DocumentType.PAPER
    assert result.document_title == "Parsed Dict Paper"
    assert len(result.hierarchy) == 1
